"""Document indexer — парсинг → chunking → embedding → skills → сохранение."""
from __future__ import annotations

import uuid
from pathlib import Path

from sqlalchemy.ext.asyncio import AsyncSession

from app.clients import embed
from app.rag.chunker import Chunk, chunk_document
from app.storage import vector_db
from app.storage.sql_db import save_chunks
from app.settings import settings as _settings


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".txt", ".html", ".xlsx"}

# Кап строк на лист xlsx — гигантские дампы (трекшены на тысячи строк) иначе
# взрывают корпус тысячами чанков-стен-из-цифр и топят нарративные доки.
_MAX_XLSX_ROWS_PER_SHEET = 300


def _strip_html(raw: str) -> str:
    """Минимальная очистка HTML — теги удаляем, &amp; и т.п. раскодируем.
    Без сторонней библиотеки чтобы не тащить bs4 ради одного формата."""
    import html as _html
    import re as _re
    text = _re.sub(r"<script\b[^>]*>.*?</script>", " ", raw, flags=_re.S | _re.I)
    text = _re.sub(r"<style\b[^>]*>.*?</style>", " ", text, flags=_re.S | _re.I)
    text = _re.sub(r"<br\s*/?>", "\n", text, flags=_re.I)
    text = _re.sub(r"</p>|</div>|</li>|</tr>", "\n", text, flags=_re.I)
    text = _re.sub(r"<[^>]+>", " ", text)
    text = _html.unescape(text)
    text = _re.sub(r"[ \t]+", " ", text)
    text = _re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _render_table_rows(rows):
    """Строки таблицы построчно в Markdown — ячейки через |. Сохраняет
    привязку метрика->значение внутри строки чанка."""
    norm = []
    for row in rows:
        cells = [" ".join(str(cell or "").split()) for cell in row]
        if not any(cells):
            continue
        norm.append("| " + " | ".join(cells) + " |")
    if not norm:
        return ""
    if len(norm) > 1:
        n_cols = norm[0].count("|") - 1
        norm.insert(1, "| " + " | ".join(["---"] * n_cols) + " |")
    return "\n".join(norm)


def _parse_docx(file_path):
    """Table-aware .docx. docx.paragraphs не включает текст таблиц — обходим
    тело документа по порядку, таблицы рендерим построчно."""
    import docx as python_docx
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table as _DocxTable
    from docx.text.paragraph import Paragraph as _DocxParagraph
    doc = python_docx.Document(file_path)
    blocks = []
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            para = _DocxParagraph(child, doc)
            t = para.text.strip()
            if not t:
                continue
            style = para.style.name if para.style else ""
            tail = style.split()[-1] if style else ""
            if style.startswith("Heading") and tail.isdigit():
                blocks.append("#" * int(tail) + " " + t)
            else:
                blocks.append(t)
        elif isinstance(child, CT_Tbl):
            rendered = _render_table_rows(
                [[cell.text for cell in row.cells] for row in _DocxTable(child, doc).rows]
            )
            if rendered:
                blocks.append(rendered)
    return "\n\n".join(blocks)


def _parse_xlsx(file_path):
    """Table-aware .xlsx через openpyxl. Каждый лист → заголовок (## имя)
    + строки в Markdown-таблице (тот же рендер, что и .docx-таблицы), чтобы
    числовой детектор видел привязку метрика->значение.

    data_only=True берёт ВЫЧИСЛЕННЫЕ значения формул (кэш, сохранённый Excel),
    а не сам текст формулы — иначе в корпус попадёт `=SUM(...)` вместо числа.
    read_only — стримим большие книги (роадмапы на мегабайты) без загрузки в RAM.
    """
    import openpyxl
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    blocks = []
    try:
        for ws in wb.worksheets:
            rows = []
            truncated = False
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= _MAX_XLSX_ROWS_PER_SHEET:
                    truncated = True
                    break
                rows.append(row)
            rendered = _render_table_rows(rows)
            if rendered:
                if truncated:
                    rendered += f"\n\n_[лист усечён до {_MAX_XLSX_ROWS_PER_SHEET} строк]_"
                blocks.append(f"## {ws.title}\n\n{rendered}")
    finally:
        wb.close()
    return "\n\n".join(blocks)


def _looks_like_garbage(text: str) -> bool:
    """Эвристика «бинарь/мусор»: на длинном тексте мало букв среди
    непробельных символов. Скан-PDF без OCR и битые конвертации `.doc` дают
    простыни спецсимволов/цифр без слов (модель их видит как «encoded
    binary»). Короткие тексты не режем — там низкая доля букв нормальна."""
    sample = text[:200_000]
    non_space = sum(1 for ch in sample if not ch.isspace())
    if non_space < 800:
        return False
    letters = sum(1 for ch in sample if ch.isalpha())
    return letters / non_space < 0.35


def parse_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            pages = [p.extract_text() or "" for p in pdf.pages]
        return "\n\n".join(pages)
    if suffix == ".docx":
        return _parse_docx(file_path)
    if suffix == ".xlsx":
        return _parse_xlsx(file_path)
    if suffix in (".md", ".txt"):
        return file_path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".html":
        return _strip_html(file_path.read_text(encoding="utf-8", errors="replace"))
    raise ValueError(f"Unsupported format: {suffix}")


def _collection_for_status(status: str) -> str:
    if status == "archived":
        return "archive"
    return "actual"  # actual, draft, unknown → в основную коллекцию


async def index_document(
    file_path: Path,
    document_id: str,
    document_metadata: dict,
    db: AsyncSession,
    workspace: str = "default",
) -> int:
    """
    Полный пайплайн индексации:
    1. Парсинг
    2. Chunking
    3. Embedding
    4. Сохранение в ChromaDB
    5. Сохранение чанков в SQLite
    6. Skills: extract_entities, track_promises
    """
    import logging as _logging
    import time as _time
    _log = _logging.getLogger(__name__)
    _t0 = _time.monotonic()

    text = parse_text(file_path)
    _log.info("index doc=%s parse %.2fs, %d chars", document_id, _time.monotonic() - _t0, len(text))

    # Guard от мусора: битые конвертации / скан-PDF без OCR парсятся в
    # «бинарь» без слов — не индексируем, чтобы не засорять корпус и не жечь
    # токены на бесполезном тексте (в этот раз чистили такое руками).
    if _looks_like_garbage(text):
        _log.warning(
            "index doc=%s: похоже на мусор/бинарь (доля букв низкая, %d симв.) — пропуск",
            document_id, len(text),
        )
        return 0

    # Кэшируем полный текст в БД — для full-mode чата без RAG-потерь.
    from app.storage.sql_db import update_document as _update_document
    try:
        await _update_document(db, document_id, {"parsed_text": text})
    except Exception:
        _log.warning("failed to cache parsed_text for doc=%s", document_id)

    _tc = _time.monotonic()
    chunks: list[Chunk] = chunk_document(text, document_metadata)
    if not chunks:
        return 0
    _log.info("index doc=%s chunk %.2fs, %d chunks", document_id, _time.monotonic() - _tc, len(chunks))

    # Батчинг эмбеддингов (не более 100 за раз)
    _te = _time.monotonic()
    batch_size = 100
    embeddings: list[list[float]] = []
    for i in range(0, len(chunks), batch_size):
        batch_texts = [c.content for c in chunks[i : i + batch_size]]
        embeddings.extend(await embed(batch_texts))
    _log.info("index doc=%s embed %.2fs", document_id, _time.monotonic() - _te)

    collection_name = _collection_for_status(document_metadata.get("status", "unknown"))

    chunk_ids = [str(uuid.uuid4()) for _ in chunks]

    # Сохраняем в ChromaDB (только не superseded)
    if document_metadata.get("status") != "superseded":
        vector_db.upsert_chunks(
            [
                {
                    "id": chunk_ids[i],
                    "content": c.content,
                    "embedding": embeddings[i],
                    "metadata": {
                        **c.metadata,
                        "document_id": document_id,
                        "title": document_metadata.get("title", ""),
                        "chunk_index": c.chunk_index,
                        "hierarchy_level": document_metadata.get("hierarchy_level", 5),
                        "status": document_metadata.get("status", "unknown"),
                    },
                }
                for i, c in enumerate(chunks)
            ],
            collection_name=collection_name,
            workspace=workspace,
        )

    # Сохраняем чанки в SQLite
    await save_chunks(
        db,
        [
            {
                "id": chunk_ids[i],
                "document_id": document_id,
                "content": c.content,
                "section": c.section,
                "subsection": c.subsection,
                "chunk_index": c.chunk_index,
                "token_count": c.token_count,
                "embedding_id": chunk_ids[i],
                "metadata_": c.metadata,
            }
            for i, c in enumerate(chunks)
        ],
    )

    # Skills pipeline — изолируем каждый skill, чтобы падение одного не
    # уносило за собой обновление chunk_count и другие skills.
    from app.skills.extract_entities import extract_and_save
    from app.skills.track_promises import extract_and_save_promises
    from app.skills.find_logic_signals import find_logic_signals_for_document
    from app.skills.find_intra_contradictions import find_intra_contradictions
    from app.skills.document_brief import generate_document_brief
    from app.storage.sql_db import update_document as _update_document

    for skill_name, coro in (_settings.ENABLE_BACKGROUND_SIGNALS and (
        ("extract_entities", extract_and_save(text, document_id, document_metadata, db)),
        ("track_promises", extract_and_save_promises(text, document_id, document_metadata, db)),
        ("find_logic_signals", find_logic_signals_for_document(document_id, db)),
        ("find_intra_contradictions", find_intra_contradictions(document_id, text, db)),
    ) or ()):
        try:
            _res = await coro
            _n = len(_res) if isinstance(_res, (list, tuple, dict)) else _res
            _log.info("skill %s doc=%s → %s", skill_name, document_id, _n)
        except Exception as exc:
            _log.exception("skill %s failed for doc=%s: %s", skill_name, document_id, exc)

    # Document-level бриф — один LLM-вызов на весь текст, сохраняем в Document.brief.
    # Это ключ к document-level анализу: full-mode подаёт LLM брифы всех документов,
    # чтобы он видел картину целиком, а не только retrieve-чанки.
    _tb = _time.monotonic()
    try:
        title = document_metadata.get("title") or "Документ"
        brief = await generate_document_brief(text, title) if _settings.ENABLE_BACKGROUND_SIGNALS else ""
        if brief:
            await _update_document(db, document_id, {"brief": brief})
            _log.info("index doc=%s brief %.2fs, %d chars",
                      document_id, _time.monotonic() - _tb, len(brief))
    except Exception as exc:
        _log.exception("brief generation failed for doc=%s: %s", document_id, exc)

    return len(chunks)


async def reindex_document(
    document_id: str,
    new_status: str,
    db: AsyncSession,
) -> None:
    """Перемещает чанки между коллекциями при смене статуса БЕЗ пересчёта
    эмбеддингов — забираем существующие векторы из Chroma и переcaем их в
    нужную коллекцию с обновлённой metadata."""
    target_collection = (
        None if new_status == "superseded" else _collection_for_status(new_status)
    )

    # Собираем все чанки документа из обеих исходных коллекций (может быть в любой).
    existing: list[dict] = []
    for src in ("actual", "archive"):
        existing.extend(vector_db.get_document_chunks(document_id, src))

    # Удаляем из обеих коллекций
    vector_db.delete_document_chunks(document_id, "actual")
    vector_db.delete_document_chunks(document_id, "archive")

    if not target_collection or not existing:
        return  # superseded не индексируется, либо чанков нет

    vector_db.upsert_chunks(
        [
            {
                "id": c["id"],
                "content": c["content"],
                "embedding": c["embedding"],
                "metadata": {**(c.get("metadata") or {}), "status": new_status},
            }
            for c in existing
            if c.get("embedding") is not None
        ],
        collection_name=target_collection,
    )
